import io
import pdfplumber
from docx import Document
import re
from app.services.resume.onet_skills import ONET_SKILLS
from app.services.resume.onet_categories import ONET_CATEGORIES


# ==========================================================
# TEXT EXTRACTION FUNCTIONS
# ==========================================================

def extract_text_from_pdf(file_path):
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])


def extract_text_from_bytes(data: bytes, ext: str) -> str:
    ext = (ext or "").lower()
    if not ext.startswith("."):
        ext = f".{ext}"

    if ext == ".pdf":
        text = ""
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text

    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join([para.text for para in doc.paragraphs])

    if ext == ".txt":
        return data.decode("utf-8", errors="ignore")

    return ""


# ==========================================================
# SECTION EXTRACTION (For Resume)
# ==========================================================

def extract_sections(text):
    sections = {
        "skills": [],
        "education": [],
        "experience": [],
        "projects": []
    }

    lines = text.split("\n")
    current_section = None

    for line in lines:
        clean_line = line.strip().lower()

        if re.search(r"\bskills\b", clean_line):
            current_section = "skills"
            continue

        elif re.search(r"\beducation\b", clean_line):
            current_section = "education"
            continue

        elif re.search(r"\bexperience\b", clean_line):
            current_section = "experience"
            continue

        elif re.search(r"\bprojects\b", clean_line):
            current_section = "projects"
            continue

        if current_section and clean_line:
            sections[current_section].append(line.strip())

    return sections


# ==========================================================
# O*NET SKILL NORMALIZATION (Robust & Case-Insensitive)
# ==========================================================
def normalize_with_onet(skill_lines):

    normalized_skills = set()

    for line in skill_lines:
        line_lower = line.lower()
        line_lower = re.sub(r"\s+", " ", line_lower)

        for canonical, data in ONET_SKILLS.items():
            aliases = data["aliases"]

            for alias in aliases:
                pattern = rf"\b{re.escape(alias.lower())}\b"

                if re.search(pattern, line_lower):
                    normalized_skills.add(canonical)
                    break

    return sorted(list(normalized_skills))


# def normalize_with_onet(skill_lines):
#     """
#     Accepts list of text blocks (resume skill section OR full JD text)
#     Returns canonical normalized skill list
#     """

#     normalized_skills = set()

#     for line in skill_lines:
#         line_lower = line.lower()

#         # Remove category labels safely
#         category_words = [
#             "programming languages",
#             "front-end technologies",
#             "backend technologies",
#             "back-end technologies",
#             "databases",
#             "tools",
#             "ai / ml",
#             "ai",
#             "ml"
#         ]

#         for word in category_words:
#             line_lower = line_lower.replace(word, "")

#         # Clean special characters
#         line_lower = re.sub(r"[•\-–]", " ", line_lower)
#         line_lower = re.sub(r"\s+", " ", line_lower)

#         # Iterate over ONET_SKILLS (word-boundary safe)
#         for key, canonical in ONET_SKILLS.items():

#             pattern = rf"\b{re.escape(key.lower())}\b"

#             if re.search(pattern, line_lower):
#                 normalized_skills.add(canonical)

#     return sorted(list(normalized_skills))


# ==========================================================
# CATEGORY MAPPING
# ==========================================================

def categorize_skills(skills):
    categorized = {}

    for category, skill_list in ONET_CATEGORIES.items():
        matched = [skill for skill in skills if skill in skill_list]
        if matched:
            categorized[category] = matched

    return categorized


# ==========================================================
# MAIN PARSER FUNCTION
# ==========================================================

def _build_resume_payload(text: str) -> dict:
    sections = extract_sections(text)
    cleaned_skills = normalize_with_onet(sections["skills"] or [text])
    categorized = categorize_skills(cleaned_skills)
    return {
        "raw_text_length": len(text),
        "sections": sections,
        "cleaned_skills": cleaned_skills,
        "onet_categorized_skills": categorized,
    }


def extract_resume_data(file_path):
    if file_path.endswith(".pdf"):
        text = extract_text_from_pdf(file_path)
    elif file_path.endswith(".docx"):
        text = extract_text_from_docx(file_path)
    elif file_path.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
            text = handle.read()
    else:
        return {"error": "Unsupported file format"}

    return _build_resume_payload(text)


def extract_resume_data_from_bytes(data: bytes, ext: str) -> dict:
    text = extract_text_from_bytes(data, ext)
    if not text.strip():
        return {"error": "Unsupported file format or empty document"}
    return _build_resume_payload(text)
